# -*- coding: utf-8 -*-
"""
从「按模块拆分测试用例」目录读取每个xlsx，一个xlsx生成一个.py测试文件

用法: python generate_test_skeleton.py
"""
import sys, os, re, glob, shutil
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')

CASES_DIR = os.path.abspath('../最终版测试用例/按模块拆分测试用例')
OUTPUT_DIR = os.path.abspath('admin-web/tests/generated')


def xlsx_to_pyname(xlsx_name):
    """01_用户端_免费资料区.xlsx → test_01_用户端_免费资料区.py"""
    name = os.path.splitext(xlsx_name)[0]
    # 去掉不合法的文件名字符
    name = re.sub(r'[（(]', '_', name)
    name = re.sub(r'[）)]', '', name)
    name = re.sub(r'[：:：]', '_', name)
    name = re.sub(r'[^\w\u4e00-\u9fff_\-]', '_', name)
    name = re.sub(r'_+', '_', name).strip('_')
    return f'test_{name}.py'


def sanitize_id(case_id):
    return case_id.replace('-', '_').replace('.', '_')


def make_class_name(xlsx_name):
    """从文件名提取类名"""
    name = os.path.splitext(xlsx_name)[0]
    # 去掉序号前缀
    name = re.sub(r'^\d+_', '', name)
    # 去掉端名前缀
    name = re.sub(r'^(用户端|商家端|管理后台)_', '', name)
    # 只保留中文字母数字
    name = re.sub(r'[^\w\u4e00-\u9fff]', '', name)
    return f'Test{name}' if name else 'TestCase'


def read_cases_from_xlsx(filepath):
    """从单个xlsx读取用例"""
    wb = openpyxl.load_workbook(filepath)
    ws = wb[wb.sheetnames[0]]
    cases = []
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(row=r, column=3).value
        if not cid or not str(cid).strip():
            continue
        cases.append({
            'id': str(cid).strip(),
            'priority': (ws.cell(row=r, column=4).value or '').strip() if ws.cell(row=r, column=4).value else '',
            'title': (ws.cell(row=r, column=6).value or '').strip(),
            'precond': (ws.cell(row=r, column=7).value or '').strip(),
            'steps': (ws.cell(row=r, column=8).value or '').strip(),
            'expected': (ws.cell(row=r, column=9).value or '').strip(),
            'note': (ws.cell(row=r, column=10).value or '').strip(),
        })
    return cases


def gen_func(c):
    """生成单条测试函数"""
    fid = sanitize_id(c['id'])
    title = c['title'].replace('"', '\\"').replace('\n', ' ')
    pri = c['priority']
    pre = c['precond'].replace('\n', '\n    #   ') if c['precond'] else '无'
    steps = c['steps'].replace('\n', '\n    #   ') if c['steps'] else '无'
    exp = c['expected'].replace('\n', '\n    #   ') if c['expected'] else '无'
    note = c['note']

    return f'''
    @case("{c['id']}", title="{title}", priority="{pri}")
    def test_{fid}(self, page):
        """[{c['id']}] {title}  [{pri}]"""
        # ── 前置条件 ──
        #   {pre}
        #
        # ── 测试步骤 ──
        #   {steps}
        #
        # ── 预期结果 ──
        #   {exp}
        #
        # ── 备注: {note} ──

        # ↓↓↓ 在这里写自动化代码，写完删掉skip ↓↓↓
        pytest.skip("待实现")
'''


def main():
    print('=' * 60)
    print('  一个xlsx → 一个.py  自动化骨架生成')
    print('=' * 60)
    print(f'  用例来源: {CASES_DIR}')
    print(f'  输出目录: {OUTPUT_DIR}')
    print()

    # 清空旧文件
    if os.path.exists(OUTPUT_DIR):
        shutil.rmtree(OUTPUT_DIR)
    os.makedirs(OUTPUT_DIR)

    xlsx_files = sorted(glob.glob(os.path.join(CASES_DIR, '*.xlsx')))
    total_files = 0
    total_cases = 0

    for xlsx_path in xlsx_files:
        fname = os.path.basename(xlsx_path)
        if fname.startswith('~$'):
            continue  # 跳过临时文件

        cases = read_cases_from_xlsx(xlsx_path)
        if not cases:
            print(f'  ⏭️  {fname} (0条，跳过)')
            continue

        py_name = xlsx_to_pyname(fname)
        py_path = os.path.join(OUTPUT_DIR, py_name)
        class_name = make_class_name(fname)

        # 按优先级排序
        pri_order = {'P0': 0, 'P1': 1, 'P2': 2, 'P3': 3, 'P4': 4, '': 5}
        cases.sort(key=lambda x: pri_order.get(x['priority'], 5))

        content = f'''# -*- coding: utf-8 -*-
"""
{os.path.splitext(fname)[0]} · 自动化测试
来源文件: {fname}
用例数量: {len(cases)} 条

运行方法:
  cd admin-web
  pytest tests/generated/{py_name} -v --headed     # 看浏览器操作
  pytest tests/generated/{py_name} -v              # 无头模式
  pytest tests/generated/{py_name} -k "test_xxx"   # 只跑某条
"""
import os
import sys
import pytest
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
from utils.case_mapping import case


class {class_name}:
    """{os.path.splitext(fname)[0]} ({len(cases)}条)"""
'''
        for c in cases:
            content += gen_func(c)

        with open(py_path, 'w', encoding='utf-8') as f:
            f.write(content)

        total_files += 1
        total_cases += len(cases)
        print(f'  ✅ {py_name:<55s} {len(cases):>3d}条  ← {fname}')

    print(f'\n{"="*60}')
    print(f'  完成: {total_files} 个文件, {total_cases} 条用例')
    print(f'  对应关系: 按模块拆分测试用例/*.xlsx → generated/test_*.py')
    print(f'{"="*60}')


if __name__ == '__main__':
    main()
