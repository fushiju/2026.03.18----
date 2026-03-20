"""
测试数据工厂

集中管理测试数据，提供各类场景的测试数据生成方法。
"""
import os
import tempfile
from decimal import Decimal


# ============ 分佣测试数据 ============

# 基础分佣场景（无分销员）
COMMISSION_BASIC_CASES = [
    {
        "id": "FC-001",
        "desc": "标准场景 - 整数",
        "paid": 100.00, "cost": 60.00,
        "platform_ratio": 50, "agent_ratio": 50,
        "expected_profit": 40.00,
        "expected_platform": 20.00,
        "expected_agent": 20.00,
    },
    {
        "id": "FC-002",
        "desc": "不同比例 - 30/70",
        "paid": 200.00, "cost": 100.00,
        "platform_ratio": 30, "agent_ratio": 70,
        "expected_profit": 100.00,
        "expected_platform": 30.00,
        "expected_agent": 70.00,
    },
    {
        "id": "FC-003",
        "desc": "利润为0",
        "paid": 50.00, "cost": 50.00,
        "platform_ratio": 50, "agent_ratio": 50,
        "expected_profit": 0.00,
        "expected_platform": 0.00,
        "expected_agent": 0.00,
    },
    {
        "id": "FC-004",
        "desc": "含小数 - 四舍五入",
        "paid": 99.99, "cost": 60.00,
        "platform_ratio": 50, "agent_ratio": 50,
        "expected_profit": 39.99,
        "expected_platform": 20.00,  # 39.99*0.5=19.995 -> 20.00
        "expected_agent": 19.99,    # 利润39.99 - 平台20.00 = 19.99（用减法避免尾差）
    },
    {
        "id": "FC-005",
        "desc": "极小利润",
        "paid": 0.01, "cost": 0.00,
        "platform_ratio": 50, "agent_ratio": 50,
        "expected_profit": 0.01,
        "expected_platform": 0.01,  # 0.01*0.5=0.005 -> 0.01
        "expected_agent": 0.00,     # 利润0.01 - 平台0.01 = 0.00（减法分配）
    },
    {
        "id": "FC-006",
        "desc": "参照今夜到家比例",
        "paid": 100.00, "cost": 0.00,
        "platform_ratio": 8, "agent_ratio": 92,
        "expected_profit": 100.00,
        "expected_platform": 8.00,
        "expected_agent": 92.00,
    },
    {
        "id": "FC-007",
        "desc": "含小数利润",
        "paid": 37.50, "cost": 20.00,
        "platform_ratio": 40, "agent_ratio": 60,
        "expected_profit": 17.50,
        "expected_platform": 7.00,
        "expected_agent": 10.50,
    },
]

# 一级分销员场景
COMMISSION_LEVEL1_CASES = [
    {
        "id": "FC-101",
        "desc": "标准一级分销",
        "paid": 100.00, "cost": 60.00,
        "platform_ratio": 50, "agent_ratio": 50, "level1_ratio": 10,
        "expected_profit": 40.00,
        "expected_platform": 20.00,
        "expected_level1": 4.00,
        "expected_agent_actual": 16.00,
    },
    {
        "id": "FC-102",
        "desc": "分销员拿走代理商全部",
        "paid": 100.00, "cost": 60.00,
        "platform_ratio": 50, "agent_ratio": 50, "level1_ratio": 50,
        "expected_profit": 40.00,
        "expected_platform": 20.00,
        "expected_level1": 20.00,
        "expected_agent_actual": 0.00,
    },
    {
        "id": "FC-103",
        "desc": "分销员比例为0",
        "paid": 100.00, "cost": 60.00,
        "platform_ratio": 50, "agent_ratio": 50, "level1_ratio": 0,
        "expected_profit": 40.00,
        "expected_platform": 20.00,
        "expected_level1": 0.00,
        "expected_agent_actual": 20.00,
    },
    {
        "id": "FC-105",
        "desc": "含小数多级计算",
        "paid": 97.20, "cost": 60.00,
        "platform_ratio": 50, "agent_ratio": 50, "level1_ratio": 10,
        "expected_profit": 37.20,
        "expected_platform": 18.60,
        "expected_level1": 3.72,
        "expected_agent_actual": 14.88,
    },
]

# 一级+二级分销员场景
COMMISSION_LEVEL2_CASES = [
    {
        "id": "FC-201",
        "desc": "标准二级分销",
        "paid": 100.00, "cost": 60.00,
        "platform_ratio": 50, "agent_ratio": 50,
        "level1_ratio": 10, "level2_ratio": 5,
        "expected_profit": 40.00,
        "expected_platform": 20.00,
        "expected_level1": 4.00,
        "expected_level2": 2.00,
        "expected_agent_actual": 14.00,
    },
    {
        "id": "FC-202",
        "desc": "一级+二级=代理商",
        "paid": 100.00, "cost": 60.00,
        "platform_ratio": 50, "agent_ratio": 50,
        "level1_ratio": 25, "level2_ratio": 25,
        "expected_profit": 40.00,
        "expected_platform": 20.00,
        "expected_level1": 10.00,
        "expected_level2": 10.00,
        "expected_agent_actual": 0.00,
    },
    {
        "id": "FC-204",
        "desc": "大代理商比例",
        "paid": 160.00, "cost": 60.00,
        "platform_ratio": 20, "agent_ratio": 80,
        "level1_ratio": 15, "level2_ratio": 10,
        "expected_profit": 100.00,
        "expected_platform": 20.00,
        "expected_level1": 15.00,
        "expected_level2": 10.00,
        "expected_agent_actual": 55.00,
    },
]


# ============ 折扣计算测试数据 ============

DISCOUNT_CALC_CASES = [
    {"id": "DC-001", "desc": "标准8折", "price": 100.00, "rate": 0.8, "expected_paid": 80.00, "expected_saving": 20.00},
    {"id": "DC-002", "desc": "8.5折含小数", "price": 99.00, "rate": 0.85, "expected_paid": 84.15, "expected_saving": 14.85},
    {"id": "DC-003", "desc": "四舍五入", "price": 33.00, "rate": 0.85, "expected_paid": 28.05, "expected_saving": 4.95},
    {"id": "DC-004", "desc": "最小金额", "price": 0.01, "rate": 0.8, "expected_paid": 0.01, "expected_saving": 0.00},
    {"id": "DC-006", "desc": "最大金额", "price": 50000.00, "rate": 0.8, "expected_paid": 40000.00, "expected_saving": 10000.00},
    {"id": "DC-008", "desc": "1折", "price": 1.00, "rate": 0.1, "expected_paid": 0.10, "expected_saving": 0.90},
    {"id": "DC-009", "desc": "9.9折", "price": 1.00, "rate": 0.99, "expected_paid": 0.99, "expected_saving": 0.01},
    {"id": "DC-010", "desc": "复杂小数", "price": 123.45, "rate": 0.73, "expected_paid": 90.12, "expected_saving": 33.33},
]


# ============ 金额输入校验测试数据 ============

# (输入值, 是否合法, 说明)
AMOUNT_INPUT_CASES = [
    ("", False, "空输入"),
    ("0", False, "零"),
    ("-1", False, "负数"),
    ("-100", False, "大负数"),
    ("0.001", False, "超精度"),
    ("0.01", True, "最小合法金额"),
    ("1", True, "正常整数"),
    ("100.50", True, "正常小数"),
    ("50000", True, "最大边界"),
    ("50000.00", True, "最大边界(小数)"),
    ("50000.01", False, "超过最大金额"),
    ("50001", False, "超过最大金额(整数)"),
    ("abc", False, "字母"),
    ("12abc34", False, "混合字母数字"),
    ("1.2.3", False, "多小数点"),
    (".", False, "纯小数点"),
    ("100.", True, "末尾小数点"),  # 可能处理为100.00
    (" 100", True, "前导空格"),   # 应trim处理
    ("1e5", False, "科学计数法"),
    ("!@#$%", False, "特殊字符"),
]


# ============ 分佣比例校验测试数据 ============

# (平台%, 代理商%, 一级%, 二级%, 是否合法, 说明)
RATIO_VALIDATION_CASES = [
    (50, 50, 0, 0, True, "标准50/50"),
    (30, 70, 0, 0, True, "30/70"),
    (8, 92, 0, 0, True, "参照今夜到家比例"),
    (0, 100, 0, 0, True, "平台0%"),
    (100, 0, 0, 0, True, "代理商0%"),
    (60, 50, 0, 0, False, "合计超100%"),
    (60, 30, 0, 0, False, "合计不足100%"),
    (-5, 105, 0, 0, False, "负数比例"),
    (50, 50, 10, 5, True, "含分销员-正常"),
    (50, 50, 25, 25, True, "分销员=代理商"),
    (50, 50, 30, 25, False, "分销员超过代理商"),
    (50, 50, 51, 0, False, "一级超过代理商"),
]


# ============ Excel导入测试数据生成 ============

def create_test_excel(filepath, rows):
    """
    生成测试用的Excel文件

    :param filepath: 输出路径
    :param rows: 数据行列表，每行是dict: {"user_id": ..., "order_amount": ..., "commission": ...}
    """
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "佣金导入"

    # 表头
    headers = ["用户ID", "订单金额", "佣金金额"]
    ws.append(headers)

    # 数据
    for row in rows:
        ws.append([row.get("user_id", ""), row.get("order_amount", ""), row.get("commission", "")])

    wb.save(filepath)
    return filepath


def create_valid_excel(filepath):
    """生成正常的Excel测试文件"""
    rows = [
        {"user_id": "U001", "order_amount": 100.00, "commission": 10.00},
        {"user_id": "U002", "order_amount": 200.00, "commission": 20.00},
        {"user_id": "U003", "order_amount": 50.00, "commission": 5.00},
    ]
    return create_test_excel(filepath, rows)


def create_invalid_excel_missing_column(filepath):
    """生成缺少必填列的Excel"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["用户ID", "订单金额"])  # 缺少"佣金金额"列
    ws.append(["U001", 100.00])
    wb.save(filepath)
    return filepath


def create_invalid_excel_negative(filepath):
    """生成含负数佣金的Excel"""
    rows = [
        {"user_id": "U001", "order_amount": 100.00, "commission": 10.00},
        {"user_id": "U002", "order_amount": 200.00, "commission": -5.00},  # 负数
    ]
    return create_test_excel(filepath, rows)


def create_invalid_excel_over_amount(filepath):
    """生成佣金大于订单金额的Excel"""
    rows = [
        {"user_id": "U001", "order_amount": 100.00, "commission": 150.00},  # 佣金>订单
    ]
    return create_test_excel(filepath, rows)


def create_empty_excel(filepath):
    """生成空Excel文件"""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.save(filepath)
    return filepath


# ============ 凭证上传测试文件生成 ============

def create_test_image(filepath, size_kb=500):
    """生成指定大小的测试图片"""
    # 生成一个简单的JPEG文件
    data = b"\xff\xd8\xff\xe0" + b"\x00" * (size_kb * 1024 - 4) + b"\xff\xd9"
    with open(filepath, "wb") as f:
        f.write(data)
    return filepath


def create_oversized_image(filepath):
    """生成超过10MB的图片"""
    return create_test_image(filepath, size_kb=10 * 1024 + 100)  # 10MB+


def create_fake_image_pdf(filepath):
    """生成伪装成jpg的PDF文件"""
    with open(filepath, "wb") as f:
        f.write(b"%PDF-1.4 fake content")
    return filepath
